"""LangChain-backed textbook loading and semantic chunking.

Supported formats:
  .pdf    — PyPDF
  .txt    — TextLoader
  .md     — TextLoader
  .docx   — Docx2txtLoader (requires docx2txt)
  .csv    — CSVLoader
  .html   — BSHTMLLoader (requires beautifulsoup4)
"""

import hashlib
import re
from pathlib import Path
from typing import List

from langchain_community.document_loaders import (
    PyPDFLoader,
    TextLoader,
    CSVLoader,
    BSHTMLLoader,
)
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from citadel.core.config import get_settings
from citadel.core.exceptions import DocumentIngestionError, EmptyDatasetError
from citadel.core.logger import get_logger
from citadel.domain.interfaces import BaseDocumentProcessor


_CHAPTER_HEADING = re.compile(
    r"^\s*(?:\d+\s+)?CHAPTER\s+(\d+)\.\s*([^\r\n]+)",
    re.IGNORECASE | re.MULTILINE,
)

_SECTION_HEADING = re.compile(
    r"^\s*(\d+(?:\.\d+)+)\s+([A-Z][^\r\n]{3,60})\s*$",
    re.MULTILINE,
)


class TextbookDocumentProcessor(BaseDocumentProcessor):
    """Load PDFs, DOCX, text, CSV, and HTML files into deterministic LangChain documents."""

    _SUPPORTED_SUFFIXES = {".pdf", ".txt", ".md", ".docx", ".csv", ".html", ".htm"}

    def __init__(
        self,
        chunk_size: int | None = None,
        chunk_overlap: int | None = None,
    ) -> None:
        """Initialize the processor with validated splitter settings."""
        settings = get_settings()
        self._chunk_size = (
            chunk_size if chunk_size is not None else settings.chunk_size
        )
        self._chunk_overlap = (
            chunk_overlap
            if chunk_overlap is not None
            else settings.chunk_overlap
        )
        if self._chunk_size <= 0:
            raise ValueError("chunk_size must be greater than zero.")
        if self._chunk_overlap < 0:
            raise ValueError("chunk_overlap cannot be negative.")
        if self._chunk_overlap >= self._chunk_size:
            raise ValueError("chunk_overlap must be less than chunk_size.")

        self._logger = get_logger(__name__)
        self._splitter = RecursiveCharacterTextSplitter(
            chunk_size=self._chunk_size,
            chunk_overlap=self._chunk_overlap,
            separators=["\n\n", "\n", ". ", " ", ""],
            add_start_index=True,
        )

    def load_and_split(self, directory_path: Path) -> List[Document]:
        """Load all supported files in a directory and return their chunks.

        Args:
            directory_path: Directory containing supported files.

        Raises:
            DocumentIngestionError: If the directory does not exist or a file
                cannot be processed.
            EmptyDatasetError: If the directory contains no supported files or
                all supported files produce no content.
        """
        if not directory_path.exists() or not directory_path.is_dir():
            raise DocumentIngestionError(
                f"Textbook directory does not exist: {directory_path}"
            )

        files = sorted(
            path
            for path in directory_path.iterdir()
            if path.is_file() and path.suffix.lower() in self._SUPPORTED_SUFFIXES
        )
        self._logger.info("Scanned %d textbook files in %s", len(files), directory_path)
        if not files:
            raise EmptyDatasetError(
                f"No supported textbook files found in {directory_path}."
            )

        chunks: List[Document] = []
        for file_path in files:
            chunks.extend(self.process_single_file(file_path))

        if not chunks:
            raise EmptyDatasetError(
                f"Supported textbook files in {directory_path} contained no text."
            )

        self._logger.info("Produced %d total textbook chunks", len(chunks))
        return chunks

    def process_single_file(
        self,
        file_path: Path,
        source_name: str | None = None,
        document_id: str | None = None,
    ) -> List[Document]:
        """Load and split one supported file into LangChain document chunks."""
        if not file_path.is_file() or file_path.suffix.lower() not in self._SUPPORTED_SUFFIXES:
            raise DocumentIngestionError(
                f"Unsupported textbook file: {file_path}"
            )

        try:
            documents = self._load_documents(file_path)
            documents = self._add_chapter_metadata(documents)
            normalized_documents = [
                self._normalize_document(document) for document in documents
            ]
            normalized_documents = [
                document
                for document in normalized_documents
                if document.page_content.strip()
            ]
            split_documents = self._splitter.split_documents(normalized_documents)
            chunks = self._to_langchain_chunks(
                file_path,
                split_documents,
                source_name=source_name,
                document_id=document_id,
            )
            self._logger.info(
                "Processed %s: %d source pages/documents, %d chunks",
                file_path.name,
                len(documents),
                len(chunks),
            )
            return chunks
        except DocumentIngestionError:
            raise
        except Exception as exc:
            raise DocumentIngestionError(
                f"Failed to process textbook file {file_path}: {exc}"
            ) from exc

    @staticmethod
    def _add_chapter_metadata(documents: List[Document]) -> List[Document]:
        """Propagate chapter and section headings to following pages."""
        annotated_documents: List[Document] = []
        chapter_number: int | None = None
        chapter_title: str | None = None
        section_number: str | None = None
        section_title: str | None = None

        for document in documents:
            # Detect chapter headings
            match = _CHAPTER_HEADING.search(document.page_content)
            if match:
                chapter_number = int(match.group(1))
                chapter_title = match.group(2).strip().title()
                section_number = None
                section_title = None

            # Detect section headings (e.g. "2.3 The OSI Model")
            sec_match = _SECTION_HEADING.search(document.page_content)
            if sec_match:
                section_number = sec_match.group(1)
                section_title = sec_match.group(2).strip()

            metadata = dict(document.metadata)
            if chapter_number is not None and chapter_title is not None:
                metadata.update(
                    {
                        "chapter_number": chapter_number,
                        "chapter_title": chapter_title,
                    }
                )
            if section_number is not None and section_title is not None:
                metadata.update(
                    {
                        "section_number": section_number,
                        "section_title": section_title,
                    }
                )
            annotated_documents.append(
                Document(page_content=document.page_content, metadata=metadata)
            )

        return annotated_documents

    def _load_documents(self, file_path: Path) -> List[Document]:
        """Load LangChain documents using the appropriate loader for the file type."""
        suffix = file_path.suffix.lower()

        if suffix == ".pdf":
            return PyPDFLoader(str(file_path)).load()

        if suffix == ".docx":
            return self._load_docx(file_path)

        if suffix == ".csv":
            try:
                return CSVLoader(str(file_path), encoding="utf-8").load()
            except Exception:
                return CSVLoader(str(file_path), encoding="latin-1").load()

        if suffix in {".html", ".htm"}:
            return BSHTMLLoader(str(file_path), open_encoding="utf-8").load()

        # .txt and .md
        try:
            return TextLoader(str(file_path), encoding="utf-8").load()
        except UnicodeDecodeError:
            return TextLoader(str(file_path), encoding="latin-1").load()

    @staticmethod
    def _load_docx(file_path: Path) -> List[Document]:
        """Load a DOCX file, falling back gracefully if docx2txt is unavailable."""
        try:
            from langchain_community.document_loaders import Docx2txtLoader
            return Docx2txtLoader(str(file_path)).load()
        except ImportError:
            # Fallback: extract with python-docx directly
            try:
                import docx  # python-docx
                doc = docx.Document(str(file_path))
                full_text = "\n".join(
                    para.text for para in doc.paragraphs if para.text.strip()
                )
                return [Document(page_content=full_text, metadata={"source": str(file_path)})]
            except ImportError as exc:
                raise DocumentIngestionError(
                    "DOCX support requires 'python-docx' or 'docx2txt'. "
                    "Run: pip install python-docx"
                ) from exc

    @staticmethod
    def _normalize_document(document: Document) -> Document:
        """Normalize extraction artifacts while preserving document metadata."""
        content = document.page_content.replace("\x00", "")
        content = content.replace("\r\n", "\n").replace("\r", "\n")
        content = re.sub(r"\n{3,}", "\n\n", content)
        return Document(page_content=content.strip(), metadata=dict(document.metadata))

    def _to_langchain_chunks(
        self,
        file_path: Path,
        documents: List[Document],
        source_name: str | None = None,
        document_id: str | None = None,
    ) -> List[Document]:
        """Add deterministic IDs and domain metadata to LangChain documents."""
        chunks: List[Document] = []
        source_file = source_name or file_path.name
        for chunk_index, document in enumerate(documents):
            content = document.page_content.strip()
            if not content:
                continue
            raw_page = document.metadata.get("page")
            page_number = raw_page + 1 if isinstance(raw_page, int) else None
            chunk_id = self._build_chunk_id(
                source_file, page_number, chunk_index, content
            )
            metadata = dict(document.metadata)
            metadata.update(
                {
                    "chunk_id": chunk_id,
                    "source_file": source_file,
                    "document_id": document_id or source_file,
                    "page_number": page_number,
                    "chunk_index": chunk_index,
                    "total_char_length": len(content),
                }
            )
            chunks.append(Document(page_content=content, metadata=metadata))
        return chunks

    @staticmethod
    def _build_chunk_id(
        source_file: str,
        page_number: int | None,
        chunk_index: int,
        content: str,
    ) -> str:
        """Build an idempotent SHA-256 identifier for a chunk."""
        identity = f"{source_file}\0{page_number}\0{chunk_index}\0{content}"
        return hashlib.sha256(identity.encode("utf-8")).hexdigest()

    @classmethod
    def is_supported(cls, filename: str) -> bool:
        """Return True if the filename has a supported extension."""
        return Path(filename).suffix.lower() in cls._SUPPORTED_SUFFIXES
